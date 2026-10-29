import io
import os
import sys
import unittest

os.environ['EXP_SECRET_KEY'] = 'teste'
os.environ['EXP_LOGIN_USER'] = 'admin'
from werkzeug.security import generate_password_hash
os.environ['EXP_LOGIN_PASSWORD_HASH'] = generate_password_hash('senha-de-teste')
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from openpyxl import Workbook, load_workbook

from unittest import mock

import app as app_mod
import envios as e
import fontes

CAB = ['Código ML', 'Código universal', 'SKU', 'Agrupador de variações', 'Produto', 'Variaçao', 'Número de produto', 'Anúncios com este produto',
       'Tamanho', 'Volume do produto (cm3)', 'Produto Estrela\nSão os...', 'Vendas\núlt. 30 dias (R$)', 'Unidades vendidas\núlt. 30 dias',
       'Aptas para\nvenda', 'Unidades com recebimento pendente no Full', 'Unidades mínimas\npara oferecer', 'Unidades sugeridas para enviar\nComo é calculado?',
       'Prazo limite de envio', 'Urgência de envio', 'Observações']


def produto(cod, sku, mlbu, vendas, aptas, pend, sug='-', obs='Estamos trabalhando para te dar uma recomendação sobre este produto.', mlbs='MLB1, MLB2'):
    return [None, cod, 'N/A', sku, '-', 'Produto ' + sku, 'N/A', mlbu, mlbs, 'Médio', 100.0, 'Não', 'R$ 10,00', vendas, aptas, pend, 5, sug, '-', 'Urgente', obs]


def relatorio(linhas):
    wb = Workbook(); ws = wb.active
    ws['B2'] = 'Relatório de Planejamento de envios'
    for j, c in enumerate(CAB, start=2):
        ws.cell(row=9, column=j, value=c)
    for i, l in enumerate(linhas, start=10):
        for j, v in enumerate(l[1:], start=2):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def base_xlsx(itens, aba='anuncios'):
    wb = Workbook(); ws = wb.active; ws.title = aba
    ws.append(['Meses', '>>>', 1.5]); ws.append(['MLB', 'SKU', 'Estoque\nPara\nEnviar', 'Situação', 'Data', 'Vendas', 'Antes', 'Estoque Full', 'Para enviar', 'Frete', 'Qualidade', 'Custo'])
    for mlbu, sku, frete, custo in itens:
        ws.append([mlbu, sku, None, 'Clássico', None, 0, 0, 0, 0, frete, 'ok', custo])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def decidir(vendas, aptas, pend, frete=10, custo=100, **kw):
    item = {'mlbu': 'MLBU1', 'sku': 'S1', 'vendas': vendas, 'aptas': aptas, 'pendentes': pend, 'sugeridas': kw.get('sug'), 'obs': kw.get('obs', ''),
            'codigo_ml': 'X', 'produto': 'p', 'mlbs': [], 'urgencia': '', 'prazo': ''}
    base = {'X': {'frete': frete, 'custo': custo, 'erro': kw.get('erro', '')}} if not kw.get('sem_base') else {}
    return e.decidir(item, base, e.validar_parametros(kw.get('params')))


class TestRegras(unittest.TestCase):
    def test_reposicao_como_na_planilha(self):      # 17 vendidas x 1,5 - 14 = 11,5 -> 12
        r = decidir(17, 14, 0)
        self.assertEqual((r['qtd'], r['status'], r['tipo']), (12, 'enviar', 'Reposição'))

    def test_recebimento_pendente_conta_como_estoque(self):   # 21 x 1,5 = 31,5 - 53 - 43 < 0
        self.assertEqual(decidir(21, 43, 53)['status'], 'nao_precisa')

    def test_sem_estoque_e_sem_venda_envia_teste_de_3(self):
        r = decidir(0, 0, 0)
        self.assertEqual((r['qtd'], r['tipo'], r['status']), (3, 'Teste', 'enviar'))

    def test_venda_zero_em_30_dias_e_teste_mesmo_com_estoque(self):
        r = decidir(0, 5, 0)
        self.assertEqual((r['qtd'], r['tipo'], r['status']), (3, 'Teste', 'enviar'))

    def test_sem_estoque_com_poucas_vendas_nunca_menos_que_o_teste(self):
        self.assertEqual(decidir(1, 0, 0)['qtd'], 3)

    def test_sem_estoque_com_muitas_vendas_usa_a_necessidade(self):
        self.assertEqual(decidir(6, 0, 0)['qtd'], 9)

    def test_sugestao_do_ml_de_1_unidade_vira_3(self):
        r = decidir(0, 0, 0, sug=1)
        self.assertEqual(r['qtd'], 3)
        r = decidir(2, 5, 0, sug=1)       # tem estoque, mas o ML pediu só 1: arredonda para o teste
        self.assertEqual((r['qtd'], r['tipo']), (3, 'Teste'))

    def test_sugestao_maior_prevalece(self):
        self.assertEqual(decidir(0, 0, 0, sug=10)['qtd'], 10)

    def test_frete_igual_ou_acima_do_limite_barra(self):
        self.assertEqual(decidir(17, 14, 0, frete=44.9, custo=100)['status'], 'enviar')
        self.assertEqual(decidir(17, 14, 0, frete=45, custo=100)['status'], 'frete_alto')
        r = decidir(17, 14, 0, frete=71.5, custo=100)
        self.assertEqual(r['impacto_pct'], 71.5)

    def test_limite_configuravel(self):
        self.assertEqual(decidir(17, 14, 0, frete=50, custo=100, params={'limite_frete_pct': '60'})['status'], 'enviar')

    def test_meses_configuravel(self):
        self.assertEqual(decidir(17, 14, 0, params={'meses': '2'})['qtd'], 20)

    def test_dado_ausente_nunca_vira_zero(self):
        self.assertEqual(decidir(17, 14, 0, sem_base=True)['status'], 'sem_dados')
        self.assertEqual(decidir(17, 14, 0, frete=None)['status'], 'sem_dados')
        self.assertEqual(decidir(17, 14, 0, custo=0)['status'], 'sem_dados')
        self.assertEqual(decidir(17, 14, 0, custo=None)['status'], 'sem_dados')

    def test_ml_impede_envio(self):
        for obs in ('Você não pode enviar estoque porque este produto está restrito por seu peso', 'Os dados são necessários para começar a oferecer o Full.\nAdicionar dados fiscais',
                    'Era igual a outro anúncio.\nIr para anúncio ativo'):
            self.assertEqual(decidir(0, 0, 0, obs=obs)['status'], 'impedido')

    def test_voce_deixou_de_oferecer_continua_elegivel(self):
        self.assertEqual(decidir(0, 0, 0, obs='Você deixou de oferecer o Full.')['status'], 'enviar')

    def test_parametros_invalidos(self):
        for k, v in (('meses', '-1'), ('meses', 'abc'), ('limite_frete_pct', '0'), ('qtd_teste', '0'), ('qtd_teste', '999')):
            with self.assertRaises(e.ArquivoInvalido):
                e.validar_parametros({k: v})

    def test_resumo_por_sku_soma_so_o_marcado(self):
        linhas = [{'sku': 'A', 'qtd': 3}, {'sku': 'A', 'qtd': 12}, {'sku': 'B', 'qtd': 3}]
        self.assertEqual(e.resumo_sku(linhas), [{'sku': 'A', 'qtd': 15, 'produtos': 2}, {'sku': 'B', 'qtd': 3, 'produtos': 1}])


class TestArquivos(unittest.TestCase):
    def test_le_relatorio_do_ml(self):
        arq = relatorio([produto('AAA', 'S1', 'MLBU1', 17, 14, 0, mlbs='MLB1, MLB2'), produto('BBB', 'MLB5024247836', '0', 0, 0, 0, mlbs='5024247836'),
                         produto('CCC', 'S3', 'MLBU3', 2, 0, 0, sug=1)])
        itens = e.ler_relatorio(arq)
        self.assertEqual(len(itens), 3)
        self.assertEqual((itens[0]['mlbu'], itens[0]['vendas'], itens[0]['aptas'], itens[0]['mlbs']), ('MLBU1', 17, 14, ['MLB1', 'MLB2']))
        self.assertEqual((itens[1]['mlbu'], itens[1]['mlbs']), ('', ['MLB5024247836']))     # anúncio sem número de produto
        self.assertEqual((itens[0]['sugeridas'], itens[2]['sugeridas']), (None, 1))

    def test_relatorio_errado_e_duplicado(self):
        with self.assertRaises(e.ArquivoInvalido):
            e.ler_relatorio(base_xlsx([('MLBU1', 'S1', 1, 2)]))
        with self.assertRaises(e.ArquivoInvalido):
            e.ler_relatorio(relatorio([produto('AAA', 'S1', 'MLBU1', 1, 1, 1), produto('AAA', 'S2', 'MLBU2', 1, 1, 1)]))

    def test_numero_em_formato_brasileiro(self):
        self.assertEqual(e._numero('1.234,56'), e.D('1234.56'))
        self.assertIsNone(e._numero('-')); self.assertIsNone(e._numero('N/A')); self.assertIsNone(e._numero('SKU NÃO ENCONTRADO'))


class TestSite(unittest.TestCase):
    def setUp(self):
        app_mod.USUARIO, app_mod.SENHA_HASH = 'admin', os.environ['EXP_LOGIN_PASSWORD_HASH']
        app_mod.app.config['SECRET_KEY'] = 'teste'; app_mod.app.config['TESTING'] = True
        self.c = app_mod.app.test_client()

    def entrar(self):
        self.assertEqual(self.c.post('/login', data={'usuario': 'admin', 'senha': 'senha-de-teste'}).status_code, 302)

    def test_exige_login(self):
        self.assertEqual(self.c.get('/envios').status_code, 302)
        for rota in ('/api/envios/ler', '/api/envios/custos', '/api/envios/fretes', '/api/envios/calcular', '/api/envios/exportar'):
            self.assertEqual(self.c.post(rota).status_code, 401)
        self.assertEqual(self.c.get('/saude').status_code, 200)
        self.assertEqual(self.c.post('/login', data={'usuario': 'admin', 'senha': 'errada'}).status_code, 401)

    def test_origem_externa_recusada(self):
        self.entrar()
        self.assertEqual(self.c.post('/api/envios/ler', headers={'Origin': 'https://mau.example'}).status_code, 403)

    def test_fluxo_completo(self):
        self.entrar()
        self.assertEqual(self.c.get('/envios').status_code, 200)
        rel = relatorio([produto('AAA', 'S1', 'MLBU1', 17, 14, 0), produto('BBB', 'S2', 'MLBU2', 17, 14, 0),
                         produto('CCC', 'S3', 'MLBU3', 0, 0, 0), produto('DDD', 'S4', 'MLBU4', 5, 0, 0), produto('EEE', 'S3', 'MLBU5', 21, 43, 53),
                         produto('FFF', 'S6', 'MLBU6', 0, 0, 0, obs='Os dados são necessários para começar a oferecer o Full.')])
        import json
        d = self.c.post('/api/envios/ler', data={'relatorio': (rel, 'rel.xlsx'), 'meses': '1.5', 'limite_frete_pct': '45', 'qtd_teste': '3'}).get_json()
        self.assertTrue(d['ok'])
        # só buscam frete/custo os que precisam de envio e que o ML não impede (EEE já tem estoque, FFF é impedido)
        self.assertEqual([p['codigo_ml'] for p in d['pendencias']], ['AAA', 'BBB', 'CCC', 'DDD'])
        self.assertEqual(d['skus'], ['S1', 'S2', 'S3', 'S4'])
        base = {'AAA': {'frete': 7.15, 'custo': 10}, 'BBB': {'frete': 30, 'custo': 100}, 'CCC': {'frete': 5, 'custo': 100},
                'DDD': {'frete': None, 'custo': 20, 'erro': 'MLB1: Anúncio não ativo'}}
        rel2 = relatorio([produto('AAA', 'S1', 'MLBU1', 17, 14, 0), produto('BBB', 'S2', 'MLBU2', 17, 14, 0),
                          produto('CCC', 'S3', 'MLBU3', 0, 0, 0), produto('DDD', 'S4', 'MLBU4', 5, 0, 0), produto('EEE', 'S3', 'MLBU5', 21, 43, 53)])
        d = self.c.post('/api/envios/calcular', data={'relatorio': (rel2, 'rel.xlsx'), 'base': json.dumps(base), 'meses': '1.5', 'limite_frete_pct': '45', 'qtd_teste': '3'}).get_json()
        self.assertTrue(d['ok'])
        st = {l['codigo_ml']: (l['status'], l['qtd']) for l in d['linhas']}
        self.assertEqual(st['AAA'], ('frete_alto', 12))          # 7,15 / 10 = 71,5%
        self.assertEqual(st['BBB'], ('enviar', 12))              # 30 / 100 = 30%
        self.assertEqual(st['CCC'], ('enviar', 3))
        self.assertEqual(st['DDD'], ('sem_dados', 8))
        self.assertIn('Anúncio não ativo', {l['codigo_ml']: l['motivo'] for l in d['linhas']}['DDD'])
        self.assertEqual(st['EEE'][0], 'nao_precisa')

    def test_candidatos_endpoints(self):
        self.entrar()
        self.assertEqual(self.c.get('/candidatos').status_code, 200)
        with mock.patch.object(fontes, 'pedidos_pagina', lambda conta, de, ate, off: {'itens': [], 'total': 0, 'proximo': 0, 'concluido': True}):
            d = self.c.post('/api/candidatos/pedidos', json={'conta': 'ML44', 'de': '2026-10-01', 'ate': '2026-10-03', 'offset': 0}).get_json()
        self.assertTrue(d['ok'])
        for ruim in ({'conta': 'XX', 'de': '2026-10-01', 'ate': '2026-10-03'}, {'conta': 'ML19', 'de': 'ontem', 'ate': '2026-10-03'},
                     {'conta': 'ML19', 'de': '2026-10-05', 'ate': '2026-10-03'}, {'conta': 'ML19', 'de': '2026-01-01', 'ate': '2026-10-03'},
                     {'conta': 'ML19', 'de': '2026-10-01', 'ate': '2026-10-03', 'offset': -1}):
            self.assertEqual(self.c.post('/api/candidatos/pedidos', json=ruim).status_code, 400)
        with mock.patch.object(fontes, 'itens_info', lambda conta, ids: {i: {} for i in ids}):
            self.assertTrue(self.c.post('/api/candidatos/itens', json={'conta': 'ML19', 'ids': ['MLB123456789']}).get_json()['ok'])
        self.assertEqual(self.c.post('/api/candidatos/itens', json={'conta': 'ML19', 'ids': ['http://x']}).status_code, 400)
        linha = {'conta': 'ML19', 'sku': '=CMD()', 'produto': 'P', 'mlbs': 'MLB1', 'envio': 'Flex', 'situacao': 'Ativo', 'unidades': 3, 'pedidos': 2, 'valor': 30, 'preco_medio': 10, 'por_dia': 0.1, 'custo': 4, 'custo_total': 12, 'estoque': 5}
        r = self.c.post('/api/candidatos/exportar', json={'linhas': [linha], 'periodo': 'x'})
        self.assertEqual(r.status_code, 200)
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(ws['B2'].value, "'=CMD()"); self.assertEqual(ws['G2'].value, 3)
        self.assertEqual(self.c.post('/api/candidatos/exportar', json={'linhas': []}).status_code, 400)

    def test_custos_endpoint(self):
        self.entrar()
        with mock.patch.object(fontes, 'custos', lambda skus: {s: 10.0 for s in skus}):
            d = self.c.post('/api/envios/custos', json={'skus': ['A', 'B']}).get_json()
        self.assertEqual(d['custos'], {'A': 10.0, 'B': 10.0})
        self.assertEqual(self.c.post('/api/envios/custos', json={'skus': []}).status_code, 400)
        self.assertEqual(self.c.post('/api/envios/custos', json={'skus': [1]}).status_code, 400)
        with mock.patch.object(fontes, 'custos', side_effect=fontes.FonteIndisponivel('fora do ar')):
            r = self.c.post('/api/envios/custos', json={'skus': ['A']})
        self.assertEqual((r.status_code, r.get_json()['error']), (503, 'fora do ar'))

    def test_fretes_endpoint_valida_e_repassa(self):
        self.entrar()
        chamadas = []
        def falso(conta, itens, espera_max=24):
            chamadas.append((conta, itens)); return {i['codigo_ml']: {'frete': 9.75, 'erro': '', 'mlb': 'MLB1'} for i in itens}, []
        with mock.patch.object(fontes, 'fretes', falso):
            d = self.c.post('/api/envios/fretes', json={'conta': 'ML19', 'itens': [{'codigo_ml': 'AAA', 'mlbs': ['MLB123']}]}).get_json()
        self.assertEqual(d['fretes']['AAA']['frete'], 9.75); self.assertEqual(chamadas[0][0], 'ML19')
        for ruim in ({'conta': 'XX', 'itens': [{'codigo_ml': 'A', 'mlbs': ['MLB1']}]}, {'conta': 'ML19', 'itens': []},
                     {'conta': 'ML19', 'itens': [{'codigo_ml': 'A', 'mlbs': ['http://x']}]},
                     {'conta': 'ML19', 'itens': [{'codigo_ml': 'A', 'mlbs': ['MLB1']}] * 11}):
            self.assertEqual(self.c.post('/api/envios/fretes', json=ruim).status_code, 400)

    def test_calcular_valida_arquivo_e_base(self):
        self.entrar()
        r = self.c.post('/api/envios/calcular', data={'relatorio': (io.BytesIO(b'texto'), 'rel.xlsx'), 'base': '{"a":{}}'})
        self.assertEqual(r.status_code, 400)
        r = self.c.post('/api/envios/calcular', data={'relatorio': (relatorio([produto('A', 'S', 'M', 1, 1, 1)]), 'rel.csv'), 'base': '{"a":{}}'})
        self.assertEqual(r.status_code, 400)
        r = self.c.post('/api/envios/calcular', data={'relatorio': (relatorio([produto('A', 'S', 'M', 1, 1, 1)]), 'rel.xlsx'), 'base': '[1]'})
        self.assertEqual(r.status_code, 400)

    def linhas_exp(self):
        base = {'status': 'enviar', 'tipo': 'Reposição', 'codigo_ml': 'A', 'mlbu': 'MLBU1', 'sku': 'S1', 'produto': 'Prod', 'mlbs': ['MLB1', 'MLB2'],
                'motivo': 'm', 'vendas': 1, 'estoque_total': 0, 'frete': 1, 'custo': 10, 'impacto_pct': 10, 'urgencia': 'Urgente', 'prazo': '01/11'}
        return [dict(base, incluir=True, qtd=12, qtd_calc=12),
                dict(base, codigo_ml='B', mlbu='MLBU2', sku='S1', incluir=True, qtd=5, qtd_calc=3, tipo='Teste'),
                dict(base, codigo_ml='C', mlbu='MLBU3', sku='=CMD()', produto='=HYPERLINK("x")', incluir=True, qtd=4, qtd_calc=4, status='frete_alto'),
                dict(base, codigo_ml='D', mlbu='MLBU4', sku='S9', incluir=False, qtd=7, qtd_calc=7)]

    def test_exportar_envio_e_producao(self):
        self.entrar()
        r = self.c.post('/api/envios/exportar', json={'tipo': 'envio', 'linhas': self.linhas_exp(), 'parametros': {'meses': '1.5'}, 'arquivo': 'rel.xlsx'})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ['Enviar', 'Resumo por SKU', 'Fora do envio', 'Parâmetros'])
        env = [[c.value for c in l] for l in wb['Enviar'].iter_rows(min_row=2)]
        self.assertEqual(len(env), 4)                                   # 3 marcados + TOTAL
        self.assertEqual(env[-1][0], 'TOTAL'); self.assertEqual(env[-1][5], 21)
        origens = {l[0]: l[7] for l in env[:-1]}
        self.assertEqual(origens['A'], 'Regra')
        self.assertIn('Quantidade editada', origens['B'])
        self.assertIn('manualmente', origens['C'])
        self.assertTrue(any(str(l[2]).startswith("'=") for l in env))     # texto com = fica inofensivo
        self.assertEqual({l[0] for l in wb['Fora do envio'].iter_rows(min_row=2, values_only=True)}, {'Enviar'})
        resumo = {l[0]: l[1] for l in wb['Resumo por SKU'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(resumo['S1'], 17); self.assertEqual(resumo['TOTAL'], 21); self.assertNotIn('S9', resumo)
        r2 = self.c.post('/api/envios/exportar', json={'tipo': 'producao', 'linhas': self.linhas_exp()})
        self.assertEqual(load_workbook(io.BytesIO(r2.data)).sheetnames, ['Resumo por SKU'])

    def test_exportar_valida(self):
        self.entrar()
        self.assertEqual(self.c.post('/api/envios/exportar', json={'tipo': 'x', 'linhas': self.linhas_exp()}).status_code, 400)
        self.assertEqual(self.c.post('/api/envios/exportar', json={'tipo': 'envio', 'linhas': []}).status_code, 400)
        nada = [dict(l, incluir=False) for l in self.linhas_exp()]
        self.assertEqual(self.c.post('/api/envios/exportar', json={'tipo': 'envio', 'linhas': nada}).status_code, 400)
        ruim = self.linhas_exp(); ruim[0]['qtd'] = 10 ** 7
        self.assertEqual(self.c.post('/api/envios/exportar', json={'tipo': 'envio', 'linhas': ruim}).status_code, 400)


if __name__ == '__main__':
    unittest.main()
