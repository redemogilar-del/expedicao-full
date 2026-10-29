import io
import os
import sys
import unittest
from unittest import mock

os.environ['EXP_SECRET_KEY'] = 'teste'
os.environ['EXP_LOGIN_USER'] = 'admin'
from werkzeug.security import generate_password_hash
os.environ['EXP_LOGIN_PASSWORD_HASH'] = generate_password_hash('senha-de-teste')
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import app as app_mod
import fontes
import vendas as v

P = v.validar_parametros({'imposto_pct': 8})          # imposto 8% digitado, margem padrão 20%


def venda(**kw):
    base = {'sku': 'A', 'quantidade': 2, 'venda': 100.0, 'tarifa': 15.0, 'frete': 10.0, 'grupo': 'a_enviar'}
    base.update(kw)
    return base


class TestRegraLucro(unittest.TestCase):
    def test_conta_do_lucro(self):
        r = v.analisar(venda(), {'custo': 30.0, 'margem': 10.0}, P)
        # repasse 100-15-10=75; imposto 8; custo 60 -> lucro 7; margem 7/60
        self.assertEqual((r['repasse'], r['imposto'], r['custo_total'], r['lucro'], r['margem']), (75.0, 8.0, 60.0, 7.0, 11.67))
        self.assertEqual(r['classe'], 'dentro')                # 11,67 >= 10
        self.assertEqual(v.analisar(venda(), {'custo': 30.0, 'margem': 20.0}, P)['classe'], 'abaixo')

    def test_prejuizo(self):
        r = v.analisar(venda(frete=50.0), {'custo': 30.0, 'margem': 20.0}, P)
        self.assertEqual((r['lucro'], r['classe']), (-33.0, 'prejuizo'))

    def test_margem_exatamente_na_meta_esta_dentro(self):
        r = v.analisar(venda(quantidade=1, venda=100.0, tarifa=0.0, frete=0.0), {'custo': 50.0, 'margem': 84.0}, P)   # lucro 42 / 50 = 84%
        self.assertEqual((r['margem'], r['classe']), (84.0, 'dentro'))

    def test_dado_ausente_nunca_vira_zero(self):
        for campo, cad, falta in (('frete', {'custo': 30.0}, 'frete'), ('tarifa', {'custo': 30.0}, 'tarifa'), ('venda', {'custo': 30.0}, 'valor da venda')):
            r = v.analisar(venda(**{campo: None}), cad, P)
            self.assertEqual((r['classe'], r['lucro'], r['margem']), ('sem_dados', None, None))
            self.assertIn(falta, r['pendencias'])
        self.assertIn('custo', v.analisar(venda(), {'custo': None}, P)['pendencias'])
        self.assertIn('custo', v.analisar(venda(), None, P)['pendencias'])
        self.assertIn('SKU', v.analisar(venda(sku=''), {'custo': None}, P)['pendencias'])
        self.assertEqual(v.analisar(venda(sku=''), {'custo': 30.0}, P)['classe'], 'abaixo')          # custo digitado na tela dispensa o SKU

    def test_margem_padrao_e_full(self):
        self.assertEqual((v.analisar(venda(), {'custo': 30.0}, P)['meta'], v.analisar(venda(), {'custo': 30.0}, P)['meta_origem']), (20.0, 'padrão'))
        r = v.analisar(venda(grupo='full'), {'custo': 30.0, 'margem': 10.0, 'margem_full': 5.0}, P)
        self.assertEqual((r['meta'], r['meta_origem']), (5.0, 'planilha'))

    def test_imposto_da_planilha_por_conta_e_full(self):
        pl = v.validar_parametros({'impostos': {'ML19': {'normal': '8', 'full': '5'}, 'ML44': {'normal': '10', 'full': '6'}}})
        self.assertIsNone(pl['imposto_pct'])
        cad = {'custo': 30.0, 'margem': 10.0}
        self.assertEqual(v.analisar(venda(conta='ML19'), cad, pl)['imposto'], 8.0)
        self.assertEqual(v.analisar(venda(conta='ML44'), cad, pl)['imposto'], 10.0)
        self.assertEqual(v.analisar(venda(conta='ML19', grupo='full'), cad, pl)['imposto'], 5.0)
        r = v.analisar(venda(conta='ML44', grupo='full'), cad, pl); self.assertEqual((r['imposto'], r['imposto_pct']), (6.0, 6.0))
        # o digitado na tela vale para tudo e manda sobre a planilha
        dig = v.validar_parametros({'imposto_pct': 12, 'impostos': {'ML19': {'normal': 8, 'full': 5}}})
        self.assertEqual(v.analisar(venda(conta='ML19'), cad, dig)['imposto'], 12.0)
        # sem planilha e sem digitado: sem dados, nunca zero
        r = v.analisar(venda(conta='ML19'), cad, v.validar_parametros({}))
        self.assertEqual((r['classe'], r['imposto']), ('sem_dados', None)); self.assertIn('imposto', r['pendencias'])
        with self.assertRaises(v.ParametroInvalido):
            v.validar_parametros({'impostos': 'x'})

    def test_perda_em_dinheiro(self):
        pre = v.analisar(venda(frete=50.0), {'custo': 30.0, 'margem': 20.0}, P)
        self.assertEqual((pre['perda'], pre['perda_tipo']), (33.0, 'prejuizo'))
        abx = v.analisar(venda(), {'custo': 30.0, 'margem': 20.0}, P)                  # lucro 7; para 20% de 60 = 12 -> deixou 5
        self.assertEqual((abx['perda'], abx['perda_tipo']), (5.0, 'deixou'))
        ok = v.analisar(venda(), {'custo': 30.0, 'margem': 10.0}, P)
        self.assertEqual((ok['perda'], ok['perda_tipo']), (0.0, ''))
        self.assertIsNone(v.analisar(venda(frete=None), {'custo': 30.0}, P)['perda'])
        r = v.calcular([venda(frete=50.0), venda()], {'A': {'custo': 30.0, 'margem': 20.0}}, P)[1]
        self.assertEqual((r['prejuizo_rs'], r['deixou_rs']), (33.0, 5.0))

    def test_parametros(self):
        for ruim in ({'imposto_pct': 'abc'}, {'imposto_pct': -1}, {'imposto_pct': 61}, {'margem_padrao': 'x'}):
            with self.assertRaises(v.ParametroInvalido):
                v.validar_parametros(ruim)
        self.assertEqual(float(v.validar_parametros({'imposto_pct': '7,5'})['imposto_pct']), 7.5)

    def test_resumo_ignora_sem_dados(self):
        linhas, r = v.calcular([venda(), venda(frete=50.0), venda(tarifa=None)], {'A': {'custo': 30.0, 'margem': 10.0}}, P)
        self.assertEqual((r['dentro'], r['prejuizo'], r['sem_dados'], r['lucro'], r['venda']), (1, 1, 1, -26.0, 200.0))


class Resp:
    def __init__(self, dados): self.dados = dados
    def json(self): return self.dados
    def raise_for_status(self): pass


class TestFontesVendas(unittest.TestCase):
    def test_percentual(self):
        casos = {'20%': 20.0, '20,5 %': 20.5, 0.2: 20.0, '20': 20.0, '': None, 'abc': None, None: None, '-5%': None}
        for entrada, esperado in casos.items():
            self.assertEqual(fontes._percentual(entrada), esperado, entrada)

    def test_cadastro(self):
        os.environ['EXP_SHEETS_URL'] = 'https://script.google.com/macros/s/x/exec'; os.environ['EXP_SHEETS_KEY'] = 'k'
        resp = {'ok': True, 'items': {'A': {'custo': 'R$ 12,50', 'margem': '25%', 'margem_full': ''}, 'B': {'custo': '0'}}}
        with mock.patch.object(fontes._SESS, 'post', return_value=Resp(resp)):
            r = fontes.cadastro(['A', 'B'])
        self.assertEqual(r['A'], {'custo': 12.5, 'margem': 25.0, 'margem_full': None})
        self.assertEqual(r['B'], {'custo': None, 'margem': None, 'margem_full': None})

    def test_rotulo_envio(self):
        casos = [('ready_to_ship', 'invoice_pending', 'Gerar NF'), ('ready_to_ship', 'ready_to_print', 'Gerar etiqueta'), ('ready_to_ship', 'printed', 'Etiqueta impressa'),
                 ('ready_to_ship', 'ready_for_pickup', 'Pronto para coleta'), ('shipped', '', 'Enviado'), ('delivered', '', 'Entregue'), ('cancelled', '', 'Cancelado'),
                 ('handling', '', 'Preparando'), ('algo_novo', 'x', 'algo_novo / x'), ('', '', 'Não consultado')]
        for st, sub, esperado in casos:
            self.assertEqual(fontes.rotulo_envio(st, sub), esperado)

    def test_frete_do_vendedor(self):
        self.assertEqual(fontes._frete_vendedor({'senders': [{'user_id': 1, 'cost': 12.3}, {'user_id': 99, 'cost': 5}]}, '99'), 5.0)
        self.assertEqual(fontes._frete_vendedor({'senders': [{'user_id': 99, 'cost': 0}]}, '99'), 0.0)       # frete zero informado é zero de verdade
        self.assertIsNone(fontes._frete_vendedor({'senders': []}, '99'))
        self.assertIsNone(fontes._frete_vendedor(None, '99'))

    def pagina(self, pedidos, rotas):
        def ml(tk, caminho, params=None):
            if caminho == '/orders/search':
                return {'results': pedidos, 'paging': {'total': len(pedidos)}}
            for prefixo, resp in rotas.items():
                if caminho.startswith(prefixo):
                    if isinstance(resp, Exception):
                        raise resp
                    return resp
            raise RuntimeError('rota inesperada ' + caminho)
        with mock.patch.object(fontes, 'token', lambda c: 't'), mock.patch.object(fontes, '_vendedor', lambda c, t: '99'), mock.patch.object(fontes, '_ml', ml):
            return fontes.pedidos_financeiros('ML19', '2026-10-29', '2026-10-29')

    def ped(self, pid, itens, envio=777, **kw):
        return dict({'id': pid, 'status': 'paid', 'date_created': '2026-10-29T09:00:00.000-03:00', 'shipping': {'id': envio}, 'order_items': itens}, **kw)

    def oi(self, mlb, q, preco, fee, sku='S1'):
        return {'item': {'id': mlb, 'title': 'T ' + mlb, 'seller_sku': sku}, 'quantity': q, 'unit_price': preco, 'sale_fee': fee}

    def test_pedido_normal_e_full(self):
        peds = [self.ped(1, [self.oi('MLB1000001', 2, 50.0, 8.0)], envio=1), self.ped(2, [self.oi('MLB1000002', 1, 80.0, 12.0)], envio=2)]
        rotas = {'/shipments/1/costs': {'senders': [{'user_id': 99, 'cost': 11.0}]}, '/shipments/2/costs': {'senders': [{'user_id': 99, 'cost': 20.0}]},
                 '/shipments/1': {'status': 'ready_to_ship', 'logistic_type': 'xd_drop_off'}, '/shipments/2': {'status': 'ready_to_ship', 'logistic_type': 'fulfillment'}}
        r = self.pagina(peds, rotas)
        a, b = r['itens']
        self.assertEqual((a['venda'], a['tarifa'], a['frete'], a['grupo']), (100.0, 16.0, 11.0, 'a_enviar'))     # tarifa por unidade × quantidade
        self.assertEqual((b['grupo'], b['frete']), ('full', 20.0))
        self.assertTrue(r['concluido'])

    def test_enviado_e_cancelado(self):
        peds = [self.ped(1, [self.oi('MLB1000001', 1, 50.0, 8.0)], envio=1), self.ped(2, [self.oi('MLB1000002', 1, 50.0, 8.0)], envio=2, status='cancelled')]
        rotas = {'/shipments/1/costs': {'senders': [{'user_id': 99, 'cost': 3.0}]}, '/shipments/1': {'status': 'delivered', 'logistic_type': 'xd_drop_off'}}
        itens = self.pagina(peds, rotas)['itens']
        self.assertEqual([(i['pedido'], i['grupo']) for i in itens], [('1', 'enviado')])

    def test_falha_de_frete_nao_vira_zero(self):
        peds = [self.ped(1, [self.oi('MLB1000001', 1, 50.0, 8.0)], envio=1)]
        itens = self.pagina(peds, {'/shipments/1/costs': RuntimeError('HTTP 500'), '/shipments/1': {'status': 'ready_to_ship', 'logistic_type': 'self_service'}})['itens']
        self.assertIsNone(itens[0]['frete']); self.assertIn('Frete', itens[0]['obs'])

    def test_varios_itens_dividem_o_frete_pelo_valor(self):
        peds = [self.ped(1, [self.oi('MLB1000001', 1, 30.0, 4.0), self.oi('MLB1000002', 1, 90.0, 12.0)], envio=1)]
        itens = self.pagina(peds, {'/shipments/1/costs': {'senders': [{'user_id': 99, 'cost': 20.0}]}, '/shipments/1': {'status': 'ready_to_ship', 'logistic_type': 'self_service'}})['itens']
        self.assertEqual([i['frete'] for i in itens], [5.0, 15.0])

    def test_carrinho_divide_o_frete_entre_os_pedidos(self):
        peds = [self.ped(1, [self.oi('MLB1000001', 1, 50.0, 8.0)], envio=1, pack_id=55)]
        rotas = {'/shipments/1/costs': {'senders': [{'user_id': 99, 'cost': 20.0}]}, '/shipments/1': {'status': 'ready_to_ship', 'logistic_type': 'self_service'}, '/packs/55': {'orders': [{'id': 1}, {'id': 2}]}}
        i = self.pagina(peds, rotas)['itens'][0]
        self.assertEqual(i['frete'], 10.0); self.assertIn('Carrinho', i['obs'])
        rotas['/packs/55'] = RuntimeError('HTTP 404')
        self.assertIsNone(self.pagina(peds, rotas)['itens'][0]['frete'])


class TestEndpointsVendas(unittest.TestCase):
    def setUp(self):
        app_mod.app.config['TESTING'] = True
        self.c = app_mod.app.test_client()
        self.assertEqual(self.c.post('/login', data={'usuario': 'admin', 'senha': 'senha-de-teste'}).status_code, 302)

    def linhas(self):
        base = {'conta': 'ML19', 'pedido': '1', 'data': '2026-10-29 09:00', 'mlb': 'MLB1', 'sku': 'A', 'titulo': 'P', 'grupo': 'a_enviar', 'obs': '', 'meta_origem': 'planilha',
                'quantidade': 2, 'venda': 100, 'tarifa': 15, 'frete': 10, 'custo_unit': 30, 'meta': 20, 'status_rotulo': 'Aguardando'}
        return [dict(base, nao_enviar=True), dict(base, pedido='2', frete=1, nao_enviar=False), dict(base, pedido='3', tarifa=None, sku='=CMD()')]

    def test_exige_login(self):
        c = app_mod.app.test_client()
        for rota in ('/api/vendas/pedidos', '/api/vendas/cadastro', '/api/vendas/calcular', '/api/vendas/exportar', '/api/vendas/salvar'):
            self.assertEqual(c.post(rota).status_code, 401)

    def test_pagina_e_pedidos(self):
        self.assertEqual(self.c.get('/vendas').status_code, 200)
        with mock.patch.object(fontes, 'pedidos_financeiros', lambda *a: {'itens': [], 'total': 0, 'proximo': 0, 'concluido': True}):
            self.assertTrue(self.c.post('/api/vendas/pedidos', json={'conta': 'ML19', 'de': '2026-10-29', 'ate': '2026-10-29'}).get_json()['ok'])
        for ruim in ({'conta': 'XX', 'de': '2026-10-29', 'ate': '2026-10-29'}, {'conta': 'ML19', 'de': 'hoje', 'ate': '2026-10-29'}, {'conta': 'ML19', 'de': '2026-10-30', 'ate': '2026-10-29'},
                     {'conta': 'ML19', 'de': '2026-01-01', 'ate': '2026-10-29'}, {'conta': 'ML19', 'de': '2026-10-29', 'ate': '2026-10-29', 'offset': -1}):
            self.assertEqual(self.c.post('/api/vendas/pedidos', json=ruim).status_code, 400)

    def test_calcular(self):
        raw = [{'conta': 'ML19', 'sku': 'A', 'quantidade': 2, 'venda': 100, 'tarifa': 15, 'frete': 10, 'grupo': 'a_enviar'}]
        d = self.c.post('/api/vendas/calcular', json={'linhas': raw, 'cadastro': {'A': {'custo': 30, 'margem': 10}}, 'parametros': {'imposto_pct': 8, 'margem_padrao': 20}}).get_json()
        self.assertEqual((d['linhas'][0]['lucro'], d['linhas'][0]['classe'], d['resumo']['dentro']), (7.0, 'dentro', 1))
        self.assertEqual(self.c.post('/api/vendas/calcular', json={'linhas': raw, 'cadastro': {}, 'parametros': {'imposto_pct': 'x'}}).status_code, 400)

    def test_impostos_endpoint(self):
        with mock.patch.object(fontes, 'impostos_conta', lambda c: {'normal': 8.0, 'full': 5.0} if c == 'ML19' else (_ for _ in ()).throw(fontes.FonteIndisponivel('sem TaxasML'))):
            d = self.c.post('/api/vendas/impostos', json={'contas': ['ML19', 'ML44']}).get_json()
        self.assertEqual(d['impostos'], {'ML19': {'normal': 8.0, 'full': 5.0}}); self.assertIn('ML44', d['erros'])
        self.assertEqual(self.c.post('/api/vendas/impostos', json={'contas': ['XX']}).status_code, 400)

    def test_exportar_prejuizo_ordenado_e_manual(self):
        import openpyxl
        base = self.linhas()[0]
        lin = [dict(base, pedido='1', frete=5, nao_enviar=False), dict(base, pedido='2', frete=60, nao_enviar=False), dict(base, pedido='3', frete=1, nao_enviar=False)]
        r = self.c.post('/api/vendas/exportar', json={'tipo': 'prejuizo', 'linhas': lin, 'parametros': {'imposto_pct': 8}})
        self.assertEqual(r.status_code, 200)
        dados = list(openpyxl.load_workbook(io.BytesIO(r.data))['Prejuízo e abaixo da margem'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([d[4] for d in dados], ['2'])                                  # frete 60 = prejuízo; frete 5 e 1 = dentro da margem (fora do Excel)
        lin2 = [dict(base, pedido='1', frete=50, nao_enviar=False), dict(base, pedido='2', frete=70, nao_enviar=False)]
        dados = list(openpyxl.load_workbook(io.BytesIO(self.c.post('/api/vendas/exportar', json={'tipo': 'prejuizo', 'linhas': lin2, 'parametros': {'imposto_pct': 8}}).data))['Prejuízo e abaixo da margem'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([d[4] for d in dados], ['2', '1'])                              # maior prejuízo primeiro
        lin3 = [dict(base, pedido='D', frete=5, nao_enviar=False), dict(base, pedido='A1', frete=8, nao_enviar=False), dict(base, pedido='A2', frete=14, nao_enviar=False), dict(base, pedido='P', frete=50, nao_enviar=False)]
        dados = list(openpyxl.load_workbook(io.BytesIO(self.c.post('/api/vendas/exportar', json={'tipo': 'prejuizo', 'linhas': lin3, 'parametros': {'imposto_pct': 8}}).data))['Prejuízo e abaixo da margem'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(d[1], d[4]) for d in dados], [('Prejuízo', 'P'), ('Abaixo da margem', 'A2'), ('Abaixo da margem', 'A1')])   # prejuízo primeiro, depois abaixo (maior perda antes); dentro fica de fora
        self.assertEqual(self.c.post('/api/vendas/exportar', json={'tipo': 'prejuizo', 'linhas': [dict(base, frete=5, nao_enviar=False)], 'parametros': {'imposto_pct': 8}}).status_code, 400)
        m = self.c.post('/api/vendas/exportar-manual', json={'periodo': 'x', 'itens': [{'tipo': 'Não calculado', 'conta': 'ML19', 'mlb': 'MLB1', 'sku': '=X', 'produto': 'P', 'situacao': 'Prejuízo', 'preco_atual': 10, 'preco_necessario': None, 'motivo': 'Preços diferentes por variação'}]})
        self.assertEqual(m.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(m.data))['Fazer na mão']
        self.assertEqual([c.value for c in ws[2]][:4], ['Não calculado', 'ML19', 'MLB1', "'=X"])
        self.assertEqual(self.c.post('/api/vendas/exportar-manual', json={'itens': []}).status_code, 400)

    def test_exportar_nao_enviar(self):
        import openpyxl
        r = self.c.post('/api/vendas/exportar', json={'tipo': 'nao_enviar', 'linhas': self.linhas(), 'parametros': {'imposto_pct': 8}, 'periodo': '29/10/2026'})
        self.assertEqual(r.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(r.data))['Não enviar']
        dados = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual([d[0] for d in dados], ['1'])                                   # só o marcado
        self.assertEqual(dados[0][9], 7.0); self.assertIn('abaixo do mínimo de 20.00%', dados[0][13])
        self.assertEqual(dados[0][12], 5.0)                                              # deixou de ganhar: 60 × 20% − 7 = 5
        self.assertEqual(dados[0][4], 'Aguardando')
        sem = [dict(l, nao_enviar=False) for l in self.linhas()]
        self.assertEqual(self.c.post('/api/vendas/exportar', json={'tipo': 'nao_enviar', 'linhas': sem, 'parametros': {'imposto_pct': 8}}).status_code, 400)

    def test_exportar_dia_recalcula_e_protege_formula(self):
        import openpyxl
        lin = self.linhas(); lin[0]['lucro'] = 99999                                  # lucro "pronto" vindo da tela é ignorado
        r = self.c.post('/api/vendas/exportar', json={'tipo': 'dia', 'linhas': lin, 'parametros': {'imposto_pct': 8}})
        ws = openpyxl.load_workbook(io.BytesIO(r.data))['Vendas do dia']
        dados = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual([d[1] for d in dados], ['Abaixo da margem', 'Dentro da margem', 'Sem dados']); self.assertEqual(dados[0][17], 7.0); self.assertEqual(dados[0][2], 5.0)
        self.assertEqual(dados[2][6], "'=CMD()"); self.assertEqual((dados[0][0], dados[1][0]), ('SIM', None)); self.assertEqual(dados[1][2], None)
        self.assertEqual(self.c.post('/api/vendas/exportar', json={'tipo': 'xx', 'linhas': lin, 'parametros': {'imposto_pct': 8}}).status_code, 400)

    def test_salvar(self):
        enviado = {}
        def falso(p): enviado.update(p); return {'abas': ['Vendas 2026-10-29']}
        with mock.patch.object(fontes, 'salvar_relatorio', falso):
            d = self.c.post('/api/vendas/salvar', json={'dia': '2026-10-29', 'periodo': '29/10/2026', 'linhas': self.linhas(), 'parametros': {'imposto_pct': 8}}).get_json()
        self.assertTrue(d['ok'])
        self.assertEqual((enviado['dia'], enviado['resumo']['itens'], enviado['resumo']['nao_enviar'], len(enviado['nao_enviar']), enviado['resumo']['sem_dados']), ('2026-10-29', 3, 1, 1, 1))
        chaves = [l[0] for l in enviado['base']]
        self.assertEqual(len(chaves), 3); self.assertEqual(len(set(chaves)), 3)               # cada venda tem chave própria
        self.assertEqual(enviado['base'][0][:7], ['ML19|1|MLB1|A', '2026-10-29', enviado['base'][0][2], 'ML19', '1', 'A', 'MLB1'])
        self.assertEqual(len(enviado['cab_base']), len(enviado['base'][0]))
        dup = self.linhas()[:1] * 2                                                              # mesmo item duas vezes no mesmo pedido
        with mock.patch.object(fontes, 'salvar_relatorio', falso):
            self.c.post('/api/vendas/salvar', json={'dia': '2026-10-29', 'linhas': dup, 'parametros': {'imposto_pct': 8}})
        self.assertEqual([l[0] for l in enviado['base']], ['ML19|1|MLB1|A', 'ML19|1|MLB1|A#2'])
        with mock.patch.object(fontes, 'salvar_relatorio', side_effect=fontes.FonteIndisponivel('sem planilha')):
            self.assertEqual(self.c.post('/api/vendas/salvar', json={'dia': '2026-10-29', 'linhas': self.linhas(), 'parametros': {'imposto_pct': 8}}).status_code, 503)
        self.assertEqual(self.c.post('/api/vendas/salvar', json={'dia': 'ontem', 'linhas': self.linhas(), 'parametros': {'imposto_pct': 8}}).status_code, 400)


if __name__ == '__main__':
    unittest.main()
